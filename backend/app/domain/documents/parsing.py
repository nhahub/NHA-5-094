"""Format-specific parsers that emit a normalized, source-located document view."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from importlib.metadata import PackageNotFoundError, version
from pathlib import PurePosixPath
from typing import Any, Dict, List, Optional

import openpyxl
import pymupdf
import xlrd
from docx import Document as WordDocument
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from PIL import Image

from app.core.config import settings
from app.domain.documents.upload import UploadValidationError, validate_upload


class DocumentParseError(ValueError):
    """The input passed upload validation but could not be parsed safely."""


@dataclass(frozen=True)
class ContentBlock:
    kind: str
    text: str
    source: Dict[str, Any]
    confidence: Optional[float] = None


@dataclass
class ParsedDocument:
    source_format: str
    parser_name: str
    parser_version: str
    blocks: List[ContentBlock] = field(default_factory=list)
    page_count: Optional[int] = None
    warnings: List[str] = field(default_factory=list)
    pages_requiring_ocr: List[int] = field(default_factory=list)

    @property
    def plain_text(self) -> str:
        return "\n".join(block.text for block in self.blocks if block.text)

    def to_json(self) -> str:
        return json.dumps(
            {
                "source_format": self.source_format,
                "parser_name": self.parser_name,
                "parser_version": self.parser_version,
                "page_count": self.page_count,
                "warnings": self.warnings,
                "pages_requiring_ocr": self.pages_requiring_ocr,
                "blocks": [
                    {
                        "kind": block.kind,
                        "text": block.text,
                        "source": block.source,
                        "confidence": block.confidence,
                    }
                    for block in self.blocks
                ],
            },
            ensure_ascii=False,
        )


def parse_document_content(filename: str, content: bytes) -> ParsedDocument:
    """Parse a validated upload into text/table/cell blocks with stable locators.

    OCR is identified as required for image inputs and textless PDF pages; OCR
    execution is a separate stage so native PDF text is never needlessly OCRed.
    """
    _, mime_type, extension = validate_upload(
        filename, content, None, settings.DOCUMENT_MAX_FILE_SIZE_BYTES
    )
    if extension == ".pdf":
        return _parse_pdf(content)
    if extension in (".png", ".jpg", ".jpeg"):
        return _inspect_image(content, extension)
    if extension == ".docx":
        return _parse_docx(content)
    if extension == ".xlsx":
        return _parse_xlsx(content)
    if extension == ".xls":
        return _parse_xls(content)
    if extension == ".csv":
        return _parse_csv(content)
    if extension == ".txt":
        return _parse_text(content)
    raise DocumentParseError(f"No parser is registered for {mime_type}.")


def _parser_version(package: str) -> str:
    try:
        return version(package)
    except PackageNotFoundError:
        return "unknown"


def _parse_pdf(content: bytes) -> ParsedDocument:
    result = ParsedDocument("pdf", "pymupdf", _parser_version("PyMuPDF"))
    try:
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise DocumentParseError("Password-protected PDFs are not supported.")
            result.page_count = pdf.page_count
            for page_index, page in enumerate(pdf):
                page_number = page_index + 1
                text_blocks = page.get_text("blocks", sort=True)
                page_text = "\n".join(str(block[4]).strip() for block in text_blocks if str(block[4]).strip())
                if not page_text:
                    result.pages_requiring_ocr.append(page_number)
                for block_index, block in enumerate(text_blocks):
                    text = str(block[4]).strip()
                    if text:
                        result.blocks.append(ContentBlock(
                            kind="text",
                            text=text,
                            source={
                                "page": page_number,
                                "bbox": [float(value) for value in block[:4]],
                                "block_index": block_index,
                            },
                        ))
                try:
                    tables = page.find_tables()
                    for table_index, table in enumerate(tables.tables):
                        rows = table.extract()
                        if rows:
                            result.blocks.append(ContentBlock(
                                kind="table",
                                text="\n".join("\t".join(str(cell or "") for cell in row) for row in rows),
                                source={"page": page_number, "table_index": table_index},
                            ))
                except (AttributeError, ValueError, RuntimeError):
                    result.warnings.append(f"Table extraction was unavailable on page {page_number}.")
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("PDF could not be parsed.") from exc
    return result


def _inspect_image(content: bytes, extension: str) -> ParsedDocument:
    try:
        with Image.open(io.BytesIO(content)) as image:
            width, height = image.size
            if width * height > 40_000_000:
                raise DocumentParseError("Image dimensions exceed the processing limit.")
            image.verify()
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("Image could not be decoded.") from exc
    return ParsedDocument(
        source_format=extension[1:],
        parser_name="pillow",
        parser_version=_parser_version("Pillow"),
        page_count=1,
        warnings=["OCR is required to extract text from this image."],
        pages_requiring_ocr=[1],
    )


def _parse_docx(content: bytes) -> ParsedDocument:
    # Recheck the package before the parser opens its ZIP/XML parts.
    _validate_office_archive(content, {"[Content_Types].xml", "word/document.xml"})
    try:
        doc = WordDocument(io.BytesIO(content))
        result = ParsedDocument("docx", "python-docx", _parser_version("python-docx"))
        body_index = 0
        for child in doc.element.body.iterchildren():
            body_index += 1
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if text:
                    result.blocks.append(ContentBlock(
                        kind="text",
                        text=text,
                        source={
                            "part": "word/document.xml",
                            "paragraph": body_index,
                            "style": paragraph.style.name if paragraph.style else None,
                        },
                    ))
            elif isinstance(child, CT_Tbl):
                table = WordTable(child, doc)
                for row_index, row in enumerate(table.rows, start=1):
                    for column_index, cell in enumerate(row.cells, start=1):
                        text = cell.text.strip()
                        if text:
                            result.blocks.append(ContentBlock(
                                kind="cell",
                                text=text,
                                source={
                                    "part": "word/document.xml",
                                    "table": body_index,
                                    "row": row_index,
                                    "column": column_index,
                                },
                            ))
        if doc.element.body.xpath(".//w:drawing"):
            result.warnings.append("Embedded images are retained in the source DOCX but are not OCRed by this parser.")
        for section_index, section in enumerate(doc.sections, start=1):
            for part_name, part in (("header", section.header), ("footer", section.footer)):
                for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
                    text = paragraph.text.strip()
                    if text:
                        result.blocks.append(ContentBlock(
                            kind="text",
                            text=text,
                            source={
                                "part": part_name,
                                "section": section_index,
                                "paragraph": paragraph_index,
                            },
                        ))
        return result
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("DOCX content could not be parsed.") from exc


def _validate_office_archive(content: bytes, required_parts: set[str]) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > settings.DOCUMENT_MAX_ARCHIVE_ENTRIES:
                raise DocumentParseError("Office document exceeds the permitted archive entry count.")
            names = set()
            expanded = 0
            for entry in entries:
                name = entry.filename
                raw_parts = name.rstrip("/").split("/")
                if not name or "\\" in name or PurePosixPath(name).is_absolute() or any(
                    part in ("", ".", "..") for part in raw_parts
                ):
                    raise DocumentParseError("Office document contains an unsafe archive path.")
                if entry.flag_bits & 1:
                    raise DocumentParseError("Encrypted Office documents are not supported.")
                if name in names:
                    raise DocumentParseError("Office document contains duplicate archive parts.")
                names.add(name)
                expanded += entry.file_size
                if (
                    entry.file_size > settings.DOCUMENT_MAX_ARCHIVE_ENTRY_SIZE_BYTES
                    or expanded > settings.DOCUMENT_MAX_ARCHIVE_EXPANDED_SIZE_BYTES
                ):
                    raise DocumentParseError("Office document exceeds the permitted expanded size.")
            if not required_parts.issubset(names):
                raise DocumentParseError("Office document is missing required package parts.")
            if "word/vbaProject.bin" in names:
                raise DocumentParseError("Macro-enabled Office documents are not supported.")
    except DocumentParseError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise DocumentParseError("Office document package is invalid.") from exc


def _parse_xlsx(content: bytes) -> ParsedDocument:
    _validate_office_archive(content, {"[Content_Types].xml", "xl/workbook.xml"})
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        result = ParsedDocument("xlsx", "openpyxl", _parser_version("openpyxl"))
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                result.warnings.append(f"Hidden worksheet '{worksheet.title}' was skipped.")
                continue
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        result.blocks.append(ContentBlock(
                            kind="cell",
                            text=_cell_text(cell.value),
                            source={"sheet": worksheet.title, "cell": cell.coordinate},
                        ))
        workbook.close()
        return result
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("XLSX workbook could not be parsed.") from exc


def _parse_xls(content: bytes) -> ParsedDocument:
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        result = ParsedDocument("xls", "xlrd", _parser_version("xlrd"))
        for sheet in workbook.sheets():
            for row_index in range(sheet.nrows):
                for column_index in range(sheet.ncols):
                    value = sheet.cell_value(row_index, column_index)
                    if value != "":
                        if sheet.cell_type(row_index, column_index) == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                        result.blocks.append(ContentBlock(
                            kind="cell",
                            text=str(value),
                            source={
                                "sheet": sheet.name,
                                "row": row_index + 1,
                                "column": column_index + 1,
                            },
                        ))
        workbook.release_resources()
        return result
    except Exception as exc:
        raise DocumentParseError("XLS workbook could not be parsed.") from exc


def _cell_text(value: Any) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _parse_csv(content: bytes) -> ParsedDocument:
    try:
        text = content.decode("utf-8-sig")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text), dialect)
        result = ParsedDocument("csv", "python-csv", "stdlib")
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                if value.strip():
                    result.blocks.append(ContentBlock(
                        kind="cell",
                        text=value,
                        source={"row": row_index, "column": column_index},
                    ))
        return result
    except (UnicodeDecodeError, csv.Error) as exc:
        raise DocumentParseError("CSV content could not be parsed as UTF-8.") from exc


def _parse_text(content: bytes) -> ParsedDocument:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("Text content could not be decoded as UTF-8.") from exc
    result = ParsedDocument("txt", "python-text", "stdlib")
    for line_number, line in enumerate(text.splitlines(), start=1):
        if line.strip():
            result.blocks.append(ContentBlock(
                kind="text",
                text=line,
                source={"line": line_number},
            ))
    return result
