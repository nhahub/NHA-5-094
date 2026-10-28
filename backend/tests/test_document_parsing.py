import io

import openpyxl
import pytest
import pymupdf
from docx import Document as WordDocument
from PIL import Image

from app.domain.documents.parsing import DocumentParseError, parse_document_content
from app.domain.documents import ocr as ocr_module


def _pdf(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def _docx() -> bytes:
    document = WordDocument()
    document.add_paragraph("Agreement terms", style="Heading 1")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Party"
    table.cell(0, 1).text = "Northwind Ltd"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    sheet["A1"] = "Invoice number"
    sheet["B1"] = "INV-42"
    sheet["A2"] = "Total"
    sheet["B2"] = 125.50
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "Do not parse"
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _png() -> bytes:
    image = Image.new("RGB", (32, 24), color="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_pdf_parser_preserves_page_and_bounding_box():
    parsed = parse_document_content("letter.pdf", _pdf("Visible PDF text"))

    assert parsed.source_format == "pdf"
    assert parsed.page_count == 1
    assert "Visible PDF text" in parsed.plain_text
    assert parsed.blocks[0].source["page"] == 1
    assert len(parsed.blocks[0].source["bbox"]) == 4


def test_text_and_csv_parsers_preserve_line_row_and_column_locations():
    text = parse_document_content("notes.txt", b"First line\nSecond line")
    assert [block.source["line"] for block in text.blocks] == [1, 2]

    csv_doc = parse_document_content("rows.csv", b"name;amount\nAda;12\n")
    assert [(block.source["row"], block.source["column"]) for block in csv_doc.blocks] == [
        (1, 1), (1, 2), (2, 1), (2, 2)
    ]
    assert csv_doc.blocks[-1].text == "12"


def test_docx_parser_preserves_paragraph_and_table_cell_locations():
    parsed = parse_document_content("agreement.docx", _docx())

    assert "Agreement terms" in parsed.plain_text
    assert any(
        block.text == "Northwind Ltd"
        and block.source["table"] == 2
        and block.source["row"] == 1
        and block.source["column"] == 2
        for block in parsed.blocks
    )


def test_xlsx_parser_preserves_sheet_and_cell_locations_and_skips_hidden_sheets():
    parsed = parse_document_content("invoice.xlsx", _xlsx())

    assert any(block.text == "125.5" and block.source == {"sheet": "Invoice", "cell": "B2"}
               for block in parsed.blocks)
    assert all("Do not parse" not in block.text for block in parsed.blocks)
    assert parsed.warnings == ["Hidden worksheet 'Hidden' was skipped."]


def test_image_and_scanned_pdf_are_marked_for_ocr():
    image = parse_document_content("scan.png", _png())
    assert image.pages_requiring_ocr == [1]
    assert "OCR is required" in image.warnings[0]

    scanned_pdf = pymupdf.open()
    scanned_pdf.new_page()
    scanned_bytes = scanned_pdf.tobytes()
    scanned_pdf.close()
    parsed_pdf = parse_document_content("scan.pdf", scanned_bytes)
    assert parsed_pdf.pages_requiring_ocr == [1]


def test_parser_rejects_corrupt_pdf_and_serializes_normalized_content():
    with pytest.raises(DocumentParseError):
        parse_document_content("broken.pdf", b"%PDF-1.7 invalid body")

    parsed = parse_document_content("terms.txt", b"Terms\nConditions")
    assert '"line": 2' in parsed.to_json()


def test_ocr_adapter_returns_word_confidence_and_image_pixel_location(monkeypatch):
    monkeypatch.setattr(
        ocr_module.pytesseract,
        "image_to_data",
        lambda *args, **kwargs: {
            "text": ["Total", "125.50", ""],
            "conf": ["90", "80", "-1"],
            "block_num": [1, 1, 1],
            "par_num": [1, 1, 1],
            "line_num": [1, 1, 2],
            "left": [10, 50, 0],
            "top": [20, 20, 0],
            "width": [30, 40, 0],
            "height": [12, 12, 0],
        },
    )
    blocks = ocr_module.ocr_image(_png(), page_number=3)

    assert len(blocks) == 1
    assert blocks[0].text == "Total 125.50"
    assert blocks[0].confidence == pytest.approx(0.85)
    assert blocks[0].source == {
        "page": 3,
        "line": 1,
        "bbox": [10, 20, 90, 32],
        "coordinate_system": "image-pixels-origin-top-left",
    }


def test_ocr_applies_only_to_pages_marked_by_the_pdf_parser(monkeypatch):
    document = pymupdf.open()
    document.new_page()
    content = document.tobytes()
    document.close()
    parsed = parse_document_content("scan.pdf", content)
    marker = parsed.blocks.copy()
    monkeypatch.setattr(
        ocr_module,
        "ocr_pdf_pages",
        lambda content, pages: {
            1: [
                ocr_module.ContentBlock(
                    kind="ocr_text",
                    text="scanned text",
                    source={"page": 1},
                    confidence=0.9,
                )
            ]
        },
    )
    result = ocr_module.apply_ocr("scan.pdf", content, parsed)

    assert result.pages_requiring_ocr == []
    assert result.blocks[len(marker)].text == "scanned text"
